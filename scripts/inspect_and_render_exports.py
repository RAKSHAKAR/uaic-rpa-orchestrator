import os
import sys
import json
import csv
import openpyxl
import pymupdf

sys.stdout.reconfigure(encoding="utf-8")

ARTIFACT_DIR = r"C:\Users\priyer\.gemini\antigravity-ide\brain\22dbbb76-035d-4e1c-aae4-7650e0c49771"

def inspect_all_formats(claim_id: str):
    print(f"\n==========================================")
    print(f"COMPREHENSIVE EXPORT INSPECTION FOR CLAIM: {claim_id}")
    print(f"==========================================")
    
    # 1. JSON
    json_path = f"scratch_test_{claim_id[:8]}.json"
    print(f"\n--- 1. INSPECTING JSON: {json_path} ---")
    if os.path.exists(json_path):
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        print(f"✓ Valid JSON structure!")
        print(f"  - Claim ID: {data.get('id')}")
        print(f"  - Claim Number: {data.get('claim_number')}")
        print(f"  - Insured Name: {data.get('insured_name')}")
        print(f"  - Claimant Name: {data.get('claimant_name')}")
        print(f"  - Status: {data.get('record_status')}")
        print(f"  - Total Scraped Cases: {len(data.get('court_cases', []))}")
        print(f"  - Bots array count: {len(data.get('bots', []))}")
        if data.get('bots'):
            for b in data['bots'][:3]:
                print(f"    * Bot: {b.get('name')} -> Status: {b.get('status')}, Cases: {b.get('cases_found')}")
    else:
        print(f"✗ File {json_path} not found!")

    # 2. CSV
    csv_path = f"scratch_test_{claim_id[:8]}.csv"
    print(f"\n--- 2. INSPECTING CSV: {csv_path} ---")
    if os.path.exists(csv_path):
        with open(csv_path, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)
        print(f"✓ Valid CSV format! Total lines: {len(rows)}")
        print(f"  - Columns ({len(rows[0])}): {', '.join(rows[0])}")
        for idx, row in enumerate(rows[1:], 1):
            print(f"  - Row {idx}: {row[:4]}...")
    else:
        print(f"✗ File {csv_path} not found!")

    # 3. XLSX
    xlsx_path = f"scratch_test_{claim_id[:8]}.xlsx"
    print(f"\n--- 3. INSPECTING EXCEL (XLSX): {xlsx_path} ---")
    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        print(f"✓ Valid Excel workbook! Sheets found: {wb.sheetnames}")
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            print(f"  * Sheet '{sheetname}': {ws.max_row} rows x {ws.max_column} columns")
            first_row = [str(cell.value) for cell in ws[1] if cell.value is not None]
            print(f"    Header: {first_row[:6]}")
            if ws.max_row > 1:
                sample_row = [str(cell.value) for cell in ws[2] if cell.value is not None]
                print(f"    Row 2:  {sample_row[:6]}")
    else:
        print(f"✗ File {xlsx_path} not found!")

    # 4. PDF
    pdf_path = f"scratch_test_{claim_id[:8]}.pdf"
    print(f"\n--- 4. INSPECTING & RENDERING PDF: {pdf_path} ---")
    if os.path.exists(pdf_path):
        doc = pymupdf.open(pdf_path)
        print(f"[PASS] Successfully opened PDF with PyMuPDF! Page count: {len(doc)}")
        print(f"  - Metadata: {doc.metadata}")
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            img_name = f"pdf_page_{claim_id[:8]}_{i+1}.png"
            img_path = os.path.join(ARTIFACT_DIR, img_name)
            pix.save(img_path)
            rect = page.rect
            text = page.get_text()
            print(f"  * Page {i+1}: {rect.width:.1f} x {rect.height:.1f} pt, Text length: {len(text)} chars")
            print(f"    Saved high-res render: {img_path}")
            first_lines = [line.strip() for line in text.split('\n') if line.strip()][:3]
            print(f"    Preview: {' | '.join(first_lines)}")
        doc.close()
    else:
        print(f"✗ File {pdf_path} not found!")

if __name__ == "__main__":
    cid = sys.argv[1] if len(sys.argv) > 1 else "32b240c2-8a7a-43a7-aa50-5f7358645f1c"
    inspect_all_formats(cid)
